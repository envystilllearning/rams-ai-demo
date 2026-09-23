"""DOCX tests (PRD §31): valid file, content, endpoint guards."""

import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services.document import DISCLAIMER, build_rams_docx

client = TestClient(app)

FORM = {
    "project_name": "Riverside Warehouse Fit-Out",
    "site_address": "Riverside Estate, Manchester",
    "client_name": "Riverside Property Group",
    "work_description": "Internal partition installation.",
}

GENERATED = {
    "project_summary": "Summary of the test project works.",
    "scope_of_work": "Install partitions.",
    "sequence_of_works": ["Step one", "Step two"],
    "hazards": [
        {
            "hazard": "Manual handling",
            "who_might_be_harmed": "Operatives",
            "existing_controls": "Team lifts",
            "initial_likelihood": 3,
            "initial_severity": 3,
            "initial_risk_score": 9,
            "additional_controls": "Mechanical aids",
            "residual_likelihood": 2,
            "residual_severity": 2,
            "residual_risk_score": 4,
        }
    ],
    "method_statement": {
        "preparation": "Prepare the site properly first.",
        "execution": "Execute works under supervision daily.",
        "completion": "Hand over with documentation done.",
    },
    "emergency_procedure": "Stop work and raise the alarm now.",
    "environmental_controls": "Segregate waste for recycling daily.",
    "ppe": ["Helmet", "Boots"],
}

PROFILE = {"company_name": "Northbridge Construction Ltd", "full_name": "John Smith"}


def test_docx_is_valid_zip_and_contains_content() -> None:
    data = build_rams_docx(FORM, GENERATED, PROFILE, "RAMS-TEST0001")
    assert len(data) > 5000  # real document, not empty

    zf = zipfile.ZipFile(io.BytesIO(data))
    assert "word/document.xml" in zf.namelist()

    doc = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in doc.paragraphs)
    tables_text = "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    full = text + "\n" + tables_text

    assert "Riverside Warehouse Fit-Out" in full
    assert "RAMS-TEST0001" in full
    assert "Northbridge Construction Ltd" in full
    assert "Manual handling" in full
    assert DISCLAIMER in full
    assert "Review and sign-off" in full


def test_docx_refuses_empty_hazards() -> None:
    with pytest.raises(ValueError):
        build_rams_docx(FORM, {**GENERATED, "hazards": []}, PROFILE, "RAMS-X")


def test_docx_deterministic() -> None:
    a = build_rams_docx(FORM, GENERATED, PROFILE, "RAMS-D1")
    b = build_rams_docx(FORM, GENERATED, PROFILE, "RAMS-D1")
    # Same logical content (timestamps inside zip may differ, so compare text)
    ta = "\n".join(p.text for p in Document(io.BytesIO(a)).paragraphs)
    tb = "\n".join(p.text for p in Document(io.BytesIO(b)).paragraphs)
    assert ta == tb


class FakeUser:
    id = "44444444-4444-4444-4444-444444444444"
    email = "docx@example.com"


def _override(user, rams_row):
    import app.routers.documents as doc_mod
    from app.core.deps import get_current_user

    class FakeRamsRepo:
        def get(self, uid, rid):
            return rams_row if rams_row and rid == "rams-1" else None

    class FakeProfilesRepo:
        def get(self, uid):
            return PROFILE

    import unittest.mock as mock

    app.dependency_overrides[get_current_user] = lambda: user
    rams_patcher = mock.patch.object(doc_mod, "RamsRepo", FakeRamsRepo)
    prof_patcher = mock.patch.object(doc_mod, "ProfilesRepo", FakeProfilesRepo)
    rams_patcher.start()
    prof_patcher.start()
    return rams_patcher, prof_patcher


def test_docx_endpoint_guards() -> None:
    # 401 without auth
    r = client.get("/api/rams/rams-1/documents/docx")
    assert r.status_code == 401

    from app.core.deps import get_current_user

    app.dependency_overrides[get_current_user] = lambda: FakeUser()
    try:
        # 404 for missing RAMS
        p1, p2 = _override(FakeUser(), None)
        r = client.get("/api/rams/rams-1/documents/docx")
        assert r.status_code == 404
        p1.stop()
        p2.stop()

        # 409 when content not generated yet
        p1, p2 = _override(FakeUser(), {"id": "rams-1", "input_data": FORM, "generated_data": None})
        r = client.get("/api/rams/rams-1/documents/docx")
        assert r.status_code == 409
        assert r.json()["code"] == "DOCUMENT_NOT_READY"
        p1.stop()
        p2.stop()
    finally:
        app.dependency_overrides.clear()


def test_docx_endpoint_downloads() -> None:
    from app.core.deps import get_current_user

    row = {
        "id": "rams-1",
        "document_number": "RAMS-DL0001",
        "input_data": FORM,
        "generated_data": GENERATED,
    }
    p1, p2 = _override(FakeUser(), row)
    app.dependency_overrides[get_current_user] = lambda: FakeUser()
    try:
        r = client.get("/api/rams/rams-1/documents/docx")
        assert r.status_code == 200
        assert "officedocument.wordprocessingml" in r.headers["content-type"]
        assert "RAMS-DL0001.docx" in r.headers["content-disposition"]
        zf = zipfile.ZipFile(io.BytesIO(r.content))
        assert "word/document.xml" in zf.namelist()
    finally:
        p1.stop()
        p2.stop()
        app.dependency_overrides.clear()
